from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_ROW_HEIGHT_RULE
from pathlib import Path

OUT = Path('Bharat_Navya_Product_Design_Pilot_Agreement.docx')

BLUE = '215AD6'
INK = '171A22'
MUTED = '5F687A'
LINE = 'D9DFEA'
PALE = 'EEF3FF'
PALE_GREY = 'F6F8FB'
WHITE = 'FFFFFF'

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_margins(cell, top=110, start=150, bottom=110, end=150):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(v))
        node.set(qn('w:type'), 'dxa')

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in('w:tcBorders')
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if edge not in kwargs:
            continue
        tag = 'w:{}'.format(edge)
        element = tcBorders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tcBorders.append(element)
        for key in ['val', 'sz', 'space', 'color']:
            if key in kwargs[edge]:
                element.set(qn('w:{}'.format(key)), str(kwargs[edge][key]))

def set_table_layout(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.first_child_found_in('w:tblW')
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(sum(widths)))
    tblW.set(qn('w:type'), 'dxa')
    tblInd = tblPr.first_child_found_in('w:tblInd')
    if tblInd is None:
        tblInd = OxmlElement('w:tblInd')
        tblPr.append(tblInd)
    tblInd.set(qn('w:w'), '120')
    tblInd.set(qn('w:type'), 'dxa')
    grid = table._tbl.tblGrid
    for col, width in zip(grid.gridCol_lst, widths):
        col.set(qn('w:w'), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is None:
                tcW = OxmlElement('w:tcW')
                tcPr.append(tcW)
            tcW.set(qn('w:w'), str(width))
            tcW.set(qn('w:type'), 'dxa')
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def set_font(run, size=10.5, color=INK, bold=False, italic=False):
    run.font.name = 'Aptos'
    run._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic

def style_paragraph(p, before=0, after=6, line=1.15):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line

def add_text(doc, text='', size=10.5, color=INK, bold=False, italic=False, before=0, after=6, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    style_paragraph(p, before, after)
    r = p.add_run(text)
    set_font(r, size, color, bold, italic)
    return p

def add_label_value(doc, label, value, before=0, after=3):
    p = doc.add_paragraph()
    style_paragraph(p, before, after)
    r = p.add_run(label)
    set_font(r, 10.5, INK, True)
    r = p.add_run(value)
    set_font(r, 10.5)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    style_paragraph(p, 13 if level == 1 else 8, 5)
    r = p.add_run(text)
    set_font(r, 13 if level == 1 else 11, BLUE if level == 1 else INK, True)
    p.paragraph_format.keep_with_next = True
    return p

def add_numbered_clause(doc, number, heading, body):
    p = doc.add_paragraph()
    style_paragraph(p, 7, 4, 1.15)
    p.paragraph_format.left_indent = Inches(0.0)
    p.paragraph_format.first_line_indent = Inches(0.0)
    r = p.add_run(f'{number}. ')
    set_font(r, 10.5, BLUE, True)
    r = p.add_run(heading)
    set_font(r, 10.5, INK, True)
    r = p.add_run(f' {body}')
    set_font(r, 10.5, INK)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style='List Bullet')
    style_paragraph(p, 0, 3, 1.15)
    r = p.add_run(text)
    set_font(r, 10.5)
    return p

def add_rule(doc):
    p = doc.add_paragraph()
    style_paragraph(p, 3, 7)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '8')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), LINE)
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p

def add_fill_line(doc, label):
    p = doc.add_paragraph()
    style_paragraph(p, 0, 4)
    r = p.add_run(label)
    set_font(r, 10, MUTED, True)
    r = p.add_run(' ' + '_' * 61)
    set_font(r, 10, MUTED)
    return p

def add_callout(doc, label, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_layout(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, PALE)
    set_cell_border(cell, top={'val':'single','sz':'6','color':'C6D5F7'}, bottom={'val':'single','sz':'6','color':'C6D5F7'}, left={'val':'single','sz':'6','color':'C6D5F7'}, right={'val':'single','sz':'6','color':'C6D5F7'})
    p = cell.paragraphs[0]
    style_paragraph(p, 0, 0, 1.15)
    r = p.add_run(label + ' ')
    set_font(r, 10, BLUE, True)
    r = p.add_run(body)
    set_font(r, 10, INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def add_footer(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style_paragraph(p, 0, 0)
    r = p.add_run('Bharat Navya | Initial Product Design Engagement | ')
    set_font(r, 8.5, MUTED)
    fld_char1 = OxmlElement('w:fldChar'); fld_char1.set(qn('w:fldCharType'), 'begin')
    instr_text = OxmlElement('w:instrText'); instr_text.set(qn('xml:space'), 'preserve'); instr_text.text = 'PAGE'
    fld_char2 = OxmlElement('w:fldChar'); fld_char2.set(qn('w:fldCharType'), 'end')
    run = p.add_run(); run._r.append(fld_char1); run._r.append(instr_text); run._r.append(fld_char2)

def add_header(section):
    header = section.header
    p = header.paragraphs[0]
    style_paragraph(p, 0, 0)
    r = p.add_run('BHARAT NAVYA')
    set_font(r, 9, BLUE, True)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT

def add_signature_block(doc, client_name_label, consultant_name_label):
    cell_table = doc.add_table(rows=1, cols=2)
    set_table_layout(cell_table, [4560, 4560])
    cell_table.cell(0, 0).text = ''
    cell_table.cell(0, 1).text = ''
    for cell in cell_table.row_cells(0):
        set_cell_border(cell, top={'val':'nil'}, bottom={'val':'nil'}, left={'val':'nil'}, right={'val':'nil'})
    left, right = cell_table.rows[0].cells
    lp = left.paragraphs[0]; style_paragraph(lp, 0, 3)
    r = lp.add_run(client_name_label); set_font(r, 10.5, INK, True)
    for label in ['Signature:', 'Date:']:
        p = left.add_paragraph(); style_paragraph(p, 0, 3)
        r = p.add_run(label + ' ' + '_' * 24); set_font(r, 10, INK)
    rp = right.paragraphs[0]; style_paragraph(rp, 0, 3)
    r = rp.add_run(consultant_name_label); set_font(r, 10.5, INK, True)
    for label in ['Signature:', 'Date:']:
        p = right.add_paragraph(); style_paragraph(p, 0, 3)
        r = p.add_run(label + ' ' + '_' * 27); set_font(r, 10, INK)

def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.82)
    sec.bottom_margin = Inches(0.78)
    sec.left_margin = Inches(0.82)
    sec.right_margin = Inches(0.82)
    sec.header_distance = Inches(0.32)
    sec.footer_distance = Inches(0.34)
    add_header(sec); add_footer(sec)

    normal = doc.styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)

    # Opening.
    add_text(doc, 'INDEPENDENT CONTRACTOR AGREEMENT', size=21, color=INK, bold=True, before=12, after=2)
    add_text(doc, 'Initial Product Design Engagement', size=12, color=MUTED, before=0, after=11)
    add_rule(doc)
    add_label_value(doc, 'Effective date: ', '26 August 2026')
    add_text(doc, 'This Independent Contractor Agreement (Agreement) is entered into by and between DESAI OPULENCE CONSULTANTS LLP, operating the Bharat Navya brand, acting through its Designated Partner Mahir Sarju Desai, with its current business and correspondence address at 510, 5th Floor, Turquoise 3, Near Gala Gymkhana, Bopal, Ahmedabad 380058 (Client), and Parth Akholkar, residing at 16, Vinayak Bungalows, IOC Road, Chandkheda, Ahmedabad (Consultant). The Client and Consultant are together the Parties.', after=8)

    add_callout(doc, 'INITIAL ENGAGEMENT AT A GLANCE', 'A 30-day, outcome-led product design engagement for the Bharat Navya platform. The Consultant will provide flexible product, UI/UX, prototype and advisory support; material production engineering or AI-automation work will be planned and agreed in writing before it begins.')

    add_heading(doc, '1. Engagement and role')
    add_numbered_clause(doc, '1.1', 'Appointment.', 'The Client appoints the Consultant as an Independent Product Designer & Product Advisor for the initial engagement term, and the Consultant accepts the appointment on the terms below.')
    add_numbered_clause(doc, '1.2', 'Independent contractor.', 'The Consultant is an independent contractor, not an employee, partner, agent or representative of the Client. Nothing in this Agreement creates employment, exclusivity, a joint venture, or an authority to bind either Party.')
    add_numbered_clause(doc, '1.3', 'Outcome-led engagement.', 'The Parties agree that this is an outcome-led engagement, not an hourly employment arrangement. The agreed scope, deliverables and review process govern the work.')

    add_heading(doc, '2. Term and fees')
    fee = doc.add_table(rows=1, cols=2)
    set_table_layout(fee, [2580, 6780])
    rows = [
        ('Initial engagement term', '30 calendar days from the Effective Date, unless ended under clause 7'),
        ('Initial engagement fee', 'INR 30,000 (Rupees Thirty Thousand only), inclusive of taxes'),
        ('Payment timing', 'INR 15,000 due at the start of the engagement; INR 15,000 due on the 30th calendar day after the Effective Date'),
        ('Ongoing arrangement', 'Before the initial term ends, the Parties will review scope, pace and technical involvement and may agree a future monthly arrangement in writing. No future term or fee is committed by this Agreement.'),
        ('Payment method', 'Bank account: 50100846852572 | IFSC: HDFC0001083'),
        ('GST status', 'Consultant is not GST-registered as of the Effective Date'),
    ]
    for i, (label, value) in enumerate(rows):
        if i > 0:
            fee.add_row()
        left, right = fee.rows[i].cells
        set_cell_shading(left, PALE_GREY)
        for cell in [left, right]:
            set_cell_border(cell, top={'val':'single','sz':'4','color':LINE}, bottom={'val':'single','sz':'4','color':LINE}, left={'val':'single','sz':'4','color':LINE}, right={'val':'single','sz':'4','color':LINE})
        p = left.paragraphs[0]; style_paragraph(p, 0, 0); r=p.add_run(label); set_font(r, 10, INK, True)
        p = right.paragraphs[0]; style_paragraph(p, 0, 0); r=p.add_run(value); set_font(r, 10)
    add_numbered_clause(doc, '2.1', 'Payment condition.', 'The Consultant may pause work if a payment remains unpaid for more than seven calendar days after its due date. Any revised timeline will be agreed in writing. The initial INR 15,000 payment reserves the Consultant’s availability and is non-refundable once work has begun, except as mutually agreed in writing.')
    add_numbered_clause(doc, '2.2', 'Expenses and external costs.', 'The Client is responsible for all approved third-party costs, including software subscriptions, API usage, hosting, tools, licences, travel, contractors and other external services. The Consultant will not incur such costs without the Client’s approval.')

    add_heading(doc, '3. Scope and deliverables')
    add_text(doc, 'During the initial engagement term, the Consultant will use reasonable professional effort to:', after=3)
    for item in [
        'Review and refine the existing prototype against the Bharat Navya product direction.',
        'Establish a light-first Bharat Navya visual system using the Client-supplied logo and agreed blue, white and near-black brand foundation.',
        'Define and design the Phase 1 experience for startup/company discovery and listings, recruitment discovery, lightweight phone-and-OTP account capture, and restrained future-ad placement.',
        'Research and synthesise five recruitment products into clear adoption and rejection decisions for Bharat Navya.',
        'Create relevant user flows, high-fidelity UI, clickable prototype updates, QA findings and developer handoff notes.',
        'Provide flexible product, UI/UX, prototype and limited technical input when it directly supports the experience, including feasible implementation guidance for the Client’s development team.',
    ]:
        add_bullet(doc, item)
    add_numbered_clause(doc, '3.1', 'Role flexibility and excluded work.', 'The Parties may jointly prioritise reasonable product, UI/UX, prototype and exploratory technical input as needs emerge. This does not create a commitment for the Consultant to own or deliver production backend or full-stack development, deployed AI automation, integrations, hosting, data operations, marketing execution, field work, paid advertising operations, lending, commerce, subscriptions, or ongoing engineering ownership. Any material production engineering or AI-automation work requires written agreement on the deliverable, responsibilities, timeline and any fee change before it begins.')
    add_numbered_clause(doc, '3.2', 'Client responsibilities.', 'The Client will provide timely access to brand assets, existing materials, prototype links, decision-makers and feedback. Delays in Client feedback or approvals may reasonably move delivery timelines.')

    add_heading(doc, '4. Working method and changes')
    add_numbered_clause(doc, '4.1', 'Cadence.', 'The Parties will maintain a weekly review or check-in cadence, supplemented by reasonable asynchronous feedback.')
    add_numbered_clause(doc, '4.2', 'Revisions.', 'Each major agreed design deliverable includes up to two planned revision rounds. Feedback should be consolidated by the Client where practical.')
    add_numbered_clause(doc, '4.3', 'Scope changes.', 'Material new modules, urgent work, production engineering work, or work that materially expands the initial engagement requires written approval by both Parties, including any effect on deliverables, timeline or fees.')

    add_heading(doc, '5. Intellectual property and portfolio')
    add_numbered_clause(doc, '5.1', 'Paid deliverables.', 'Upon full payment of all fees due, the Client will own the final, agreed project-specific deliverables created under this Agreement, including approved designs, prototype work and handoff materials.')
    add_numbered_clause(doc, '5.2', 'Consultant materials.', 'The Consultant retains ownership of pre-existing materials, reusable methods, processes, templates, libraries, tools, concepts and general know-how. The Client receives a perpetual right to use any incorporated Consultant materials only as part of the paid deliverables.')
    add_numbered_clause(doc, '5.3', 'Portfolio right.', 'After the work is publicly launched, or with the Client’s written approval before launch, the Consultant may display non-confidential final work in a portfolio or case study. The Consultant will not disclose confidential business information.')

    add_heading(doc, '6. Confidentiality')
    add_numbered_clause(doc, '6.1', 'Confidential information.', 'Each Party may receive non-public information about the other Party’s business, customers, plans, assets, systems or operations. Each Party will use that information only for this engagement and will protect it with reasonable care.')
    add_numbered_clause(doc, '6.2', 'Exceptions.', 'Confidential information does not include information that is publicly available without breach of this Agreement, already known lawfully, independently developed, or required to be disclosed by law.')

    add_heading(doc, '7. Ending the engagement')
    add_numbered_clause(doc, '7.1', 'Notice.', 'Either Party may end this Agreement with seven calendar days’ written notice. Neither Party is obliged to enter a future monthly arrangement after the initial term.')
    add_numbered_clause(doc, '7.2', 'Payment on termination.', 'The INR 30,000 initial engagement fee remains payable in full for the initial term, except where the Parties agree otherwise in writing. The Consultant will provide completed paid deliverables in their then-current form after amounts due are settled.')
    add_numbered_clause(doc, '7.3', 'Initial advance.', 'No refund is due for the initial INR 15,000 payment once work has begun because it reserves the Consultant’s availability. If work has not started, that amount will be refunded.')

    add_heading(doc, '8. General terms')
    add_numbered_clause(doc, '8.1', 'No guarantee.', 'The Consultant will provide services with reasonable professional skill and care, but does not guarantee product-market fit, revenue, user traction, technical performance, launch timing, third-party approvals, or business outcomes.')
    add_numbered_clause(doc, '8.2', 'Liability.', 'To the maximum extent permitted by applicable law, neither Party will be liable for indirect, incidental, special or consequential losses. The Consultant’s total liability under this Agreement will not exceed the fees actually paid to the Consultant under it.')
    add_numbered_clause(doc, '8.3', 'Entire agreement and changes.', 'This Agreement records the entire understanding of the Parties about this engagement. Changes must be in writing, including email or a mutually acknowledged message.')
    add_numbered_clause(doc, '8.4', 'Governing law.', 'This Agreement is governed by the laws of India. Courts in Ahmedabad, Gujarat will have jurisdiction over any dispute arising from it.')
    add_numbered_clause(doc, '8.5', 'Notices.', 'For this Agreement, the Consultant may be contacted at 9712783040 and the Client may be contacted at 8575800000. The Parties agree that a written notice or approval sent to these numbers by WhatsApp, or otherwise acknowledged in writing, is valid for ordinary working notices under this Agreement.')

    add_callout(doc, 'PRACTICAL NOTE', 'This is a commercial working agreement, not a substitute for advice from a qualified lawyer. Both Parties should review the completed details, especially the legal entity, tax and jurisdiction fields, before signing.')

    signature_break = doc.add_paragraph()
    signature_break.add_run().add_break(WD_BREAK.PAGE)
    add_heading(doc, '9. Signatures')
    add_text(doc, 'The Parties confirm that they have read and agreed to this Agreement.', after=8)
    add_signature_block(doc, 'Mahir Sarju Desai, Designated Partner, for Desai Opulence Consultants LLP (Bharat Navya)', 'Parth Akholkar')

    # Core properties
    props = doc.core_properties
    props.title = 'Bharat Navya Initial Product Design Engagement Agreement'
    props.subject = 'Independent Contractor Agreement — Initial 30-Day Engagement'
    props.author = 'Bharat Navya'
    props.comments = 'Fillable draft for mutual review.'
    doc.save(OUT)

if __name__ == '__main__':
    main()
